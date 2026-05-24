# core/exporters.py — Esportazione DOCX / PPTX (su file e in-memory bytes).
#
# Estratto da generatore_pro.py (salva_testo_docx / salva_slide_pptx) con in più le
# varianti `_bytes` che restituiscono `bytes` per st.download_button (Streamlit).
#
# Modulo foglia: non importa dagli altri moduli del package.

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt


# ─── COSTRUZIONE DOCUMENTO (helper condivisi) ──────────────────────────────────

def _costruisci_documento(contenuto: str, profilo: dict) -> Any:
    """Costruisce un oggetto ``Document`` python-docx dal contenuto generato.

    Effettua il parsing delle sezioni ``---TESTO ADATTATO---`` / ``---GLOSSARIO---``.
    """
    profilo = profilo or {}
    doc = Document()
    p = profilo.get("presentazione", {}) or {}
    font_nome = p.get("font", "Arial").split()[0]
    stile = doc.styles["Normal"]
    stile.font.name = font_nome
    stile.font.size = Pt(12)

    sezioni = contenuto.split("---")
    testo_corrente = ""
    glossario_corrente = ""
    in_glossario = False

    for sezione in sezioni:
        s = sezione.strip()
        if s == "TESTO ADATTATO":
            in_glossario = False
        elif s == "GLOSSARIO":
            in_glossario = True
        elif s:
            if in_glossario:
                glossario_corrente = s
            else:
                testo_corrente = s

    # Se non c'erano marker, usa l'intero contenuto come testo.
    if not testo_corrente and not glossario_corrente:
        testo_corrente = contenuto.strip()

    for riga in testo_corrente.split("\n"):
        if riga.strip():
            para = doc.add_paragraph(riga)
            para.paragraph_format.space_after = Pt(6)

    if glossario_corrente:
        doc.add_paragraph()
        heading = doc.add_paragraph("Parole chiave")
        heading.style = doc.styles["Heading 2"]
        for voce in glossario_corrente.split("\n"):
            if ":" in voce:
                termine, defin = voce.split(":", 1)
                para = doc.add_paragraph()
                run_t = para.add_run(termine.strip() + ": ")
                run_t.bold = True
                para.add_run(defin.strip())

    return doc


def _costruisci_presentazione(contenuto: str, profilo: dict) -> Any:
    """Costruisce un oggetto ``Presentation`` python-pptx dal contenuto generato.

    Effettua il parsing dei blocchi ``---SLIDE...``.
    """
    profilo = profilo or {}
    prs = Presentation()
    p = profilo.get("presentazione", {}) or {}
    font_nome = p.get("font", "Arial").split()[0]
    layout_vuoto = prs.slide_layouts[6]

    for blocco in contenuto.split("---SLIDE"):
        blocco = blocco.strip()
        if not blocco or blocco.isdigit():
            continue

        linee = [linea.strip() for linea in blocco.split("\n") if linea.strip()]
        titolo = ""
        punti: list[str] = []

        for linea in linee:
            if linea.startswith("TITOLO:"):
                titolo = linea.replace("TITOLO:", "").strip()
            elif linea.startswith("IMMAGINE:"):
                pass
            elif linea.startswith(("•", "-")):
                punti.append(linea.lstrip("•- ").strip())
            elif linea.startswith("SOTTOTITOLO:"):
                punti.insert(0, linea.replace("SOTTOTITOLO:", "").strip())

        if not titolo:
            continue

        slide = prs.slides.add_slide(layout_vuoto)
        tb = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(9), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p_titolo = tf.paragraphs[0]
        p_titolo.text = titolo
        if p_titolo.runs:
            p_titolo.runs[0].font.size = PptPt(28)
            p_titolo.runs[0].font.bold = True
            p_titolo.runs[0].font.name = font_nome

        if punti:
            tb2 = slide.shapes.add_textbox(Inches(0.3), Inches(1.5), Inches(9), Inches(5))
            tf2 = tb2.text_frame
            tf2.word_wrap = True
            for i, punto in enumerate(punti[:3]):
                para = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
                para.text = f"• {punto}"
                if para.runs:
                    para.runs[0].font.size = PptPt(18)
                    para.runs[0].font.name = font_nome

    return prs


# ─── SALVATAGGIO SU FILE (compat CLI) ──────────────────────────────────────────

def salva_testo_docx(contenuto: str, profilo: dict, percorso: str) -> None:
    """Salva il testo adattato come file DOCX su disco."""
    doc = _costruisci_documento(contenuto, profilo)
    doc.save(percorso)
    print(f"  Salvato: {percorso}")


def salva_slide_pptx(contenuto: str, profilo: dict, percorso: str) -> None:
    """Salva le slide come file PPTX su disco."""
    prs = _costruisci_presentazione(contenuto, profilo)
    prs.save(percorso)
    print(f"  Salvato: {percorso}")


# ─── EXPORT IN-MEMORY (per st.download_button) ─────────────────────────────────

def export_docx_bytes(testo: str, profilo: dict) -> bytes:
    """Restituisce il DOCX come ``bytes`` (per Streamlit, senza scrivere su disco)."""
    doc = _costruisci_documento(testo, profilo)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_pptx_bytes(slide_testo: str, profilo: dict) -> bytes:
    """Restituisce il PPTX come ``bytes`` (per Streamlit, senza scrivere su disco)."""
    prs = _costruisci_presentazione(slide_testo, profilo)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
