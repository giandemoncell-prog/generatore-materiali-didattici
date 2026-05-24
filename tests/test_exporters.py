"""Test per core.exporters (export in bytes per Streamlit)."""

from __future__ import annotations

import io

import pytest
from docx import Document
from pptx import Presentation

from core.exporters import export_docx_bytes, export_pptx_bytes

CONTENUTO_TESTO = (
    "---TESTO ADATTATO---\n"
    "La lotta per le investiture fu un conflitto tra papato e impero.\n"
    "Si concluse con il Concordato di Worms.\n"
    "---GLOSSARIO---\n"
    "investitura: nomina a una carica\n"
    "concordato: accordo tra Stato e Chiesa\n"
)

CONTENUTO_SLIDE = (
    "---SLIDE 1---\n"
    "TITOLO: La lotta per le investiture\n"
    "SOTTOTITOLO: Storia medievale\n"
    "---SLIDE 2---\n"
    "TITOLO: Cosa impariamo oggi\n"
    "- Cos'era il conflitto\n"
    "- Chi erano i protagonisti\n"
    "---SLIDE 3---\n"
    "TITOLO: Canossa\n"
    "IMMAGINE: castello sulla neve\n"
    "- Enrico IV chiede perdono\n"
)


# ── export_docx_bytes ─────────────────────────────────────────────────────────

def test_export_docx_restituisce_bytes(profilo_mario):
    out = export_docx_bytes(CONTENUTO_TESTO, profilo_mario)
    assert isinstance(out, bytes)


def test_export_docx_non_vuoto(profilo_mario):
    out = export_docx_bytes(CONTENUTO_TESTO, profilo_mario)
    assert len(out) > 0


def test_export_docx_e_un_file_docx(profilo_mario):
    out = export_docx_bytes(CONTENUTO_TESTO, profilo_mario)
    # I file OOXML iniziano con la firma ZIP "PK".
    assert out[:2] == b"PK"


def test_export_docx_contiene_paragrafi_attesi(profilo_mario):
    out = export_docx_bytes(CONTENUTO_TESTO, profilo_mario)
    doc = Document(io.BytesIO(out))
    testi = [p.text for p in doc.paragraphs]
    assert any("lotta per le investiture" in t for t in testi)
    # Il glossario produce un heading "Parole chiave".
    assert any("Parole chiave" in t for t in testi)


def test_export_docx_profilo_vuoto(profilo_mario):
    # Profilo vuoto non deve far fallire l'export.
    out = export_docx_bytes(CONTENUTO_TESTO, {})
    assert isinstance(out, bytes) and len(out) > 0


def test_export_docx_senza_marker(profilo_mario):
    # Contenuto senza marker -> usato interamente come testo.
    out = export_docx_bytes("Testo semplice senza sezioni.", profilo_mario)
    doc = Document(io.BytesIO(out))
    assert any("Testo semplice" in p.text for p in doc.paragraphs)


# ── export_pptx_bytes ─────────────────────────────────────────────────────────

def test_export_pptx_restituisce_bytes(profilo_mario):
    out = export_pptx_bytes(CONTENUTO_SLIDE, profilo_mario)
    assert isinstance(out, bytes)


def test_export_pptx_non_vuoto(profilo_mario):
    out = export_pptx_bytes(CONTENUTO_SLIDE, profilo_mario)
    assert len(out) > 0


def test_export_pptx_e_un_file_pptx(profilo_mario):
    out = export_pptx_bytes(CONTENUTO_SLIDE, profilo_mario)
    assert out[:2] == b"PK"


def test_export_pptx_ha_almeno_una_slide(profilo_mario):
    out = export_pptx_bytes(CONTENUTO_SLIDE, profilo_mario)
    prs = Presentation(io.BytesIO(out))
    assert len(prs.slides) >= 1


def test_export_pptx_conta_slide_corretto(profilo_mario):
    out = export_pptx_bytes(CONTENUTO_SLIDE, profilo_mario)
    prs = Presentation(io.BytesIO(out))
    # 3 blocchi SLIDE con titolo -> 3 slide.
    assert len(prs.slides) == 3


def test_export_pptx_profilo_vuoto():
    out = export_pptx_bytes(CONTENUTO_SLIDE, {})
    prs = Presentation(io.BytesIO(out))
    assert len(prs.slides) >= 1
